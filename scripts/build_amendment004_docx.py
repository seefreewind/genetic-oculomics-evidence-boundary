#!/usr/bin/env python3
"""Build the Amendment 004 HMG manuscript and cover-letter DOCX files."""
from pathlib import Path
import csv
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MAN_DIR = ROOT / "manuscript/hmg_amendment004"
FIG_DIR = ROOT / "figures/amendment004"


def set_font(run, size=11, bold=False, italic=False, color=(0, 0, 0)):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color)


def setup_document(doc, footer_text):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.72)
    sec.left_margin, sec.right_margin = Inches(0.82), Inches(0.82)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[style_name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(4)
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run(footer_text), 8, color=(80, 80, 80))


def add_rich_paragraph(doc, text, style=None, centered=False):
    p = doc.add_paragraph(style=style)
    if centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    # Basic Markdown emphasis plus literal inline-code markers.
    token_re = re.compile(r"(\*\*.*?\*\*|\*[^*]+\*|`.*?`)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            set_font(p.add_run(text[pos:match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            set_font(p.add_run(token[2:-2]), bold=True)
        elif token.startswith("`"):
            set_font(p.add_run(token[1:-1]), size=9.5)
        else:
            set_font(p.add_run(token[1:-1]), italic=True)
        pos = match.end()
    if pos < len(text):
        set_font(p.add_run(text[pos:]))
    return p


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_figure(doc, path, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.45))
    for element in run._element.xpath(".//wp:docPr"):
        element.set("descr", alt)
        element.set("title", path.stem)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = borders.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")


def add_tsv_table(doc, path, widths):
    with path.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.reader(stream, delimiter="\t"))
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    add_table_borders(table)
    for r_idx, row in enumerate(rows):
        tr = table.rows[0] if r_idx == 0 else table.add_row()
        tr._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c_idx, value in enumerate(row):
            cell = tr.cells[c_idx]
            cell.width = Inches(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            cell.text = value
            if r_idx == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "1F4E78")
                cell._tc.get_or_add_tcPr().append(shade)
            elif r_idx % 2 == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F2F6FA")
                cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_font(run, 7.2, bold=(r_idx == 0), color=((255, 255, 255) if r_idx == 0 else (0, 0, 0)))
    doc.add_paragraph()


def parse_markdown(doc, path, include_figures=False):
    figures = {
        1: (FIG_DIR / "Figure1_complementary_framework.png", "Complementary retinal systemic evidence framework"),
        2: (FIG_DIR / "Figure2_integrated_evidence.png", "Integrated genetic evidence after Amendment 004"),
        3: (FIG_DIR / "Figure3_local_inference_sensitivity.png", "Local genetic correlation inference sensitivity"),
    }
    inserted = set()
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            set_font(p.add_run(line[2:]), 18, bold=True)
        elif line.startswith("## "):
            keep_with_next(doc.add_heading(line[3:], level=1))
        elif line.startswith("### "):
            keep_with_next(doc.add_heading(line[4:], level=2))
        elif line.startswith("- "):
            add_rich_paragraph(doc, line[2:], style="List Bullet")
        else:
            add_rich_paragraph(doc, line)
            if include_figures:
                for number, (fig_path, alt) in figures.items():
                    if number not in inserted and re.search(rf"\bFig\. {number}(?!\d)", line):
                        add_figure(doc, fig_path, alt)
                        inserted.add(number)


def build_manuscript():
    doc = Document()
    setup_document(doc, "Human Molecular Genetics submission manuscript | Amendment 004")
    parse_markdown(doc, MAN_DIR / "HMG_MANUSCRIPT_AMENDMENT004.md", include_figures=True)
    keep_with_next(doc.add_heading("Main tables", level=1))
    keep_with_next(doc.add_heading("Table 1 Dataset level inventory of retinal and systemic GWAS resources", level=2))
    add_tsv_table(doc, ROOT / "tables/hmg_visual_final/Table1_dataset_level_GWAS_inventory.tsv", [1.25, 1.15, 1.1, 1.15, 1.0, 0.65, 0.8])
    doc.add_page_break()
    keep_with_next(doc.add_heading("Table 2 Retinal domain interpretation boundary", level=2))
    add_tsv_table(doc, ROOT / "tables/amendment004/Table2_retinal_domain_interpretation.tsv", [1.0, 1.75, 1.45, 1.65, 1.65])
    out = MAN_DIR / "HMG_MANUSCRIPT_AMENDMENT004.docx"
    doc.save(out)
    return out


def build_cover_letter():
    doc = Document()
    setup_document(doc, "Human Molecular Genetics cover letter | Amendment 004")
    parse_markdown(doc, MAN_DIR / "HMG_COVER_LETTER_AMENDMENT004.md")
    out = MAN_DIR / "HMG_COVER_LETTER_AMENDMENT004.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    print(build_manuscript())
    print(build_cover_letter())
